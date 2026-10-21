from __future__ import annotations

import csv
import io
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any


_HEADER_SCAN_ROWS = 25
_EXCLUDED_ROW = re.compile(r"\b(sub-?totals?|grand totals?|payment totals?|remittance totals?|totals?)\b", re.IGNORECASE)
_DUE_HEADER = re.compile(r"\b(due|balance|original|invoice total|gross)\b", re.IGNORECASE)
# Abbreviated forms (AR/ERP exports): "Tot Amt Pd", "Amt Pd", "Pd Amt", "Pmt Amt", "Amt Paid".
_PAID_HEADER = re.compile(
    r"\b(amount paid|paid amount|payment amount|amount applied|net paid|remit amount"
    r"|(?:tot(?:al)?\.?\s+)?amt\.?\s+(?:pd|paid)|(?:pd|paid)\.?\s+amt|pmt\.?\s+amt)\b",
    re.IGNORECASE,
)
_ID_HEADER = re.compile(
    r"\b(pro|invoice|freight bill|document (?:reference )?(?:number|no|#)|doc (?:ref|no|#))",
    re.IGNORECASE,
)
_NON_INVOICE_HEADER = re.compile(
    r"\b(po|load|bol|bill of lading|customer reference|your reference|payment reference|check|invoice total|invoice amount|amount due|balance)\b",
    re.IGNORECASE,
)


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_indexes(rows: list[list[Any]], scan_rows: int = _HEADER_SCAN_ROWS) -> tuple[int | None, int | None, int | None]:
    """(header row, PRO/invoice column, amount-paid column) by header meaning, or (None, None, None)."""
    best = None
    best_score = -1
    for row_index, row in enumerate(rows[:scan_rows]):
        headers = [_text(value) for value in row]
        invoice_indexes = [
            index
            for index, header in enumerate(headers)
            if _ID_HEADER.search(header) and not _NON_INVOICE_HEADER.search(header)
        ]
        paid_indexes = [
            index for index, header in enumerate(headers) if _PAID_HEADER.search(header)
        ]
        if not paid_indexes:
            paid_indexes = [
                index
                for index, header in enumerate(headers)
                if header
                and re.fullmatch(r"amount|paid|net", header, re.IGNORECASE)
                and not _DUE_HEADER.search(header)
            ]
        if invoice_indexes and paid_indexes:
            score = 10 * bool(_PAID_HEADER.search(headers[paid_indexes[0]])) + len(headers)
            if score > best_score:
                best = (row_index, invoice_indexes[0], paid_indexes[0])
                best_score = score
    if best:
        return best
    return None, None, None


# Shared with the HTML table parser so both detect columns by the same header rules.
find_header = _header_indexes
is_total_label = _EXCLUDED_ROW.search


def _parse_decimal_text(value: Any) -> str | None:
    text = _text(value)
    if not text or text.startswith("="):
        return None
    text = text.replace("$", "").replace(",", "").strip()
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1].strip()
    try:
        parsed = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    if not parsed.is_finite():
        return None
    return format(parsed, "f")


def _format_excel_identifier(value: Any, number_format: str | None = None) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        if number_format and re.fullmatch(r"0+", number_format):
            return str(value).zfill(len(number_format))
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return _format_excel_identifier(int(value), number_format)
    return str(value).strip()


def _normalize_rows(
    rows: list[list[Any]],
    *,
    filename: str,
    worksheet: str | None = None,
    identifier_formats: dict[tuple[int, int], str] | None = None,
    cell_errors: set[tuple[int, int]] | None = None,
) -> dict:
    from extraction.column_mapping import resolve_columns

    mapping = resolve_columns(rows)
    header_row, invoice_column, amount_column = mapping["headerRow"], mapping["identifierColumn"], mapping["amountPaidColumn"]
    warnings: list[str] = []
    if mapping["source"] == "agent":
        # The ACH Extraction Expert maps any header wording (or no header); the row checks below still apply.
        if mapping["tableKind"] != "payment_detail" or invoice_column is None or amount_column is None:
            kind = "spreadsheet_no_paid_amounts" if mapping["tableKind"] == "statement_or_open_items" else "spreadsheet_not_payment_detail"
            return {"paymentLines": [], "warnings": [kind], "columnMapping": mapping}
        if header_row is None:
            header_row = -1
        if mapping["rulesAgree"] is False:
            warnings.append("column_mapping_differs_from_rules")
    if header_row is None:
        if rows and len(rows[0]) >= 2:
            # A first row with no numbers is an unrecognized header, not a payment line.
            first = [_text(value) for value in rows[0][:2]]
            looks_like_header = all(first) and not any(_parse_decimal_text(value) is not None for value in first)
            header_row, invoice_column, amount_column = (0 if looks_like_header else -1), 0, 1
            warnings.append("header_row_not_detected; first two columns used and require review")
        else:
            return {"paymentLines": [], "warnings": ["no usable invoice/amount columns found"]}
    else:
        # A sheet whose amount-paid column is empty on every row is a statement or open-items list
        # (amounts owed), not payment detail: it contributes no payment lines.
        data_rows = [
            row for row in rows[header_row + 1 :]
            if invoice_column < len(row) and _text(row[invoice_column]) and not _EXCLUDED_ROW.search(_text(row[invoice_column]))
        ]
        if data_rows and all(amount_column >= len(row) or not _text(row[amount_column]) for row in data_rows):
            return {"paymentLines": [], "warnings": ["spreadsheet_no_paid_amounts"], "columnMapping": mapping}

    lines = []
    totals = []
    cell_errors = cell_errors or set()
    identifier_formats = identifier_formats or {}
    for row_index, row in enumerate(rows):
        if row_index <= header_row:
            continue
        if invoice_column >= len(row) or amount_column >= len(row):
            continue
        invoice_number = _format_excel_identifier(
            row[invoice_column],
            identifier_formats.get((row_index, invoice_column)),
        )
        amount_text = _text(row[amount_column])
        if not invoice_number and not amount_text:
            continue
        if _EXCLUDED_ROW.search(invoice_number):
            warnings.append(f"row_{row_index + 1}:total_or_subtotal_row_skipped")
            totals.append(_parse_decimal_text(amount_text))
            continue

        line_warnings = []
        if (row_index, invoice_column) in cell_errors or (row_index, amount_column) in cell_errors:
            line_warnings.append("spreadsheet_error_cell")
        if not invoice_number:
            line_warnings.append("missing_invoice_number")
        paid_amount = _parse_decimal_text(amount_text)
        if paid_amount is None:
            line_warnings.append("missing_or_invalid_amount_paid")

        line = {
            "proInvoiceNumber": invoice_number or None,
            "amountPaid": paid_amount,
            "source": {
                "row": row_index + 1,
                "worksheet": worksheet,
                "documentName": filename,
            },
            "confidence": None,
        }
        if line_warnings:
            line["warnings"] = line_warnings
        lines.append(line)

    # One Total row in the paid column is the sheet's own proof of its amounts. With subtotals (several
    # total rows) the check is skipped rather than guessed; reconciliation to the stated total remains.
    verified = False
    if len(totals) == 1 and totals[0] is not None and lines:
        line_sum = sum((Decimal(line["amountPaid"]) for line in lines if line["amountPaid"] is not None), Decimal("0"))
        verified = line_sum == Decimal(totals[0])
        if not verified:
            warnings.append("spreadsheet_total_mismatch")
    # The agent and the frozen rules chose different columns, and nothing in the sheet confirms either:
    # flagged so the validator sends the run to review when the email states no total either.
    if mapping.get("rulesAgree") is False and not verified:
        warnings.append("column_mapping_unverified")

    return {"paymentLines": lines, "warnings": warnings, "columnMapping": mapping, "totalVerified": verified}


def parse_csv(content: bytes, filename: str) -> dict:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return _normalize_rows(rows, filename=filename)


def parse_xlsx(content: bytes, filename: str) -> dict:
    from openpyxl import load_workbook
    from openpyxl.cell.cell import TYPE_ERROR

    workbook = load_workbook(io.BytesIO(content), data_only=False, read_only=False)
    all_lines = []
    warnings = []
    mappings = []
    for worksheet in workbook.worksheets:
        rows = []
        formats = {}
        errors = set()
        for row_index, row in enumerate(worksheet.iter_rows(), start=0):
            values = []
            for column_index, cell in enumerate(row):
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    errors.add((row_index, column_index))
                    warnings.append(f"{worksheet.title}!{cell.coordinate}:formula_unresolved")
                if cell.data_type == TYPE_ERROR:
                    errors.add((row_index, column_index))
                    warnings.append(f"{worksheet.title}!{cell.coordinate}:spreadsheet_error")
                formats[(row_index, column_index)] = cell.number_format
                values.append(value)
            rows.append(values)
        parsed = _normalize_rows(
            rows,
            filename=filename,
            worksheet=worksheet.title,
            identifier_formats=formats,
            cell_errors=errors,
        )
        all_lines.extend(parsed["paymentLines"])
        warnings.extend(parsed["warnings"])
        if parsed.get("columnMapping"):
            mappings.append({"worksheet": worksheet.title, **parsed["columnMapping"]})
    workbook.close()
    return {"paymentLines": all_lines, "warnings": warnings, "columnMappings": mappings}


def parse_xls(content: bytes, filename: str) -> dict:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=content)
    all_lines = []
    warnings = []
    mappings = []
    for worksheet in workbook.sheets():
        rows = []
        formats = {}
        errors = set()
        for row_index in range(worksheet.nrows):
            values = []
            for column_index in range(worksheet.ncols):
                cell = worksheet.cell(row_index, column_index)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_ERROR:
                    errors.add((row_index, column_index))
                    warnings.append(f"{worksheet.name}!R{row_index + 1}C{column_index + 1}:spreadsheet_error")
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                    xf = workbook.xf_list[cell.xf_index]
                    cell_format = workbook.format_map.get(xf.format_key)
                    if cell_format:
                        formats[(row_index, column_index)] = cell_format.format_str
                values.append(value)
            rows.append(values)
        parsed = _normalize_rows(
            rows,
            filename=filename,
            worksheet=worksheet.name,
            identifier_formats=formats,
            cell_errors=errors,
        )
        all_lines.extend(parsed["paymentLines"])
        warnings.extend(parsed["warnings"])
        if parsed.get("columnMapping"):
            mappings.append({"worksheet": worksheet.name, **parsed["columnMapping"]})
    return {"paymentLines": all_lines, "warnings": warnings, "columnMappings": mappings}


def parse_spreadsheet(content: bytes, filename: str) -> dict:
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        return parse_csv(content, filename)
    if extension == ".xlsx":
        return parse_xlsx(content, filename)
    if extension == ".xls":
        return parse_xls(content, filename)
    raise ValueError(f"Unsupported spreadsheet format: {extension or '(missing extension)'}")