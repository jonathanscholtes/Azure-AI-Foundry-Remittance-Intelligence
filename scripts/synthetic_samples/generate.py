"""Generate synthetic remittance emails (.msg) for testing, with expected results.

Every company, person, address, PRO number and amount is invented, so the set can be shared publicly and
used in any environment. The scenarios cover the remittance shapes the pipeline must
handle: tables in the email body, PDF / Excel / image / scanned attachments, a TIFF without a file
extension, very large tables, forwarded threads with several payments, disputes, short pays,
credits, duplicates, links to a portal, EDI addenda and a Spanish-language remittance.

Output (default samples/synthetic/, tracked in git; only _build/ is ignored):

    samples/synthetic/<id>/
        <id>.msg                the email (built with Outlook, never sent)
        expected.json           expected status, metadata and issues (schemas/golden-sample-v1.json)
        expected-lines.csv      expected payment lines: pro,amount
    samples/synthetic/README.md catalog of the scenarios
    samples/synthetic/_build/   the generated attachments, for inspection

The folders use the golden-set layout, so a sample can be copied into samples/golden/ and scored with
scripts/golden_eval.py. Labels are written unverified: the lines, totals and dates are exact by
construction, but the expected status and issues state the intended behavior - confirm them on a
first run, then `golden_eval.py verify`.

PROs are 9 digits (valid under the neutral 6-12 digit rule and under typical 9-10 digit
organization rules). The payee is the fictional "Foundry Freight Corp" - add it to RECEIVING_ORG_ALIASES when
testing a configured deployment so it is never read as the payer.

Requires Windows with Outlook installed (the .msg files are written through the Outlook object
model) and openpyxl.

    .venv\\Scripts\\python.exe scripts/synthetic_samples/generate.py [--out samples/synthetic] [--only id ...]
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import shutil
import subprocess
import sys
import zlib
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
PAYEE = "Foundry Freight Corp"
PAYEE_EMAIL = "ar@foundryfreight.example"
CENT = Decimal("0.01")


# --- data ----------------------------------------------------------------------------------------


@dataclass
class Line:
    pro: str
    amount: Decimal
    invoice_date: date
    gross: Decimal | None = None
    discount: Decimal = Decimal("0.00")
    note: str = ""

    def __post_init__(self):
        if self.gross is None:
            self.gross = self.amount + self.discount


def money(value: Decimal, symbol: str = "") -> str:
    sign = "-" if value < 0 else ""
    return f"{sign}{symbol}{abs(value):,.2f}"


def total(lines: list[Line]) -> Decimal:
    return sum((line.amount for line in lines), Decimal("0.00"))


class Data:
    """Seeded random PROs, amounts and dates - the same scenario always produces the same file."""

    def __init__(self, seed: int):
        self.rng = random.Random(seed)
        self.used: set[str] = set()

    def pro(self, leading_zero: bool = False) -> str:
        while True:
            value = (f"00{self.rng.randint(1000000, 9999999)}" if leading_zero
                     else str(self.rng.randint(100000000, 899999999)))
            if value not in self.used:
                self.used.add(value)
                return value

    def amount(self, low: int = 85, high: int = 4800) -> Decimal:
        return (Decimal(self.rng.randint(low * 100, high * 100)) / 100).quantize(CENT)

    def lines(self, count: int, paid_on: date, leading_zero: bool = False, discount: bool = False) -> list[Line]:
        result = []
        for _ in range(count):
            amount = self.amount()
            disc = (amount * Decimal("0.02")).quantize(CENT) if discount and self.rng.random() < 0.3 else Decimal("0.00")
            result.append(Line(self.pro(leading_zero), amount, paid_on - timedelta(days=self.rng.randint(12, 60)), discount=disc))
        return result

    def digits(self, count: int) -> str:
        return "".join(str(self.rng.randint(0, 9)) for _ in range(count))


# --- minimal PDF writer (text pages and image pages, no dependencies) -------------------------------

_WIDTHS = {**{d: 556 for d in "0123456789"}, ",": 278, ".": 278, "-": 333, "$": 556, " ": 278, "(": 333, ")": 333}


def _text_width(text: str, size: float) -> float:
    return sum(_WIDTHS.get(char, 556) for char in text) * size / 1000


def _pdf_string(text: str) -> bytes:
    raw = text.encode("cp1252", errors="replace")
    return b"(" + raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)") + b")"


class Pdf:
    """Letter-size pages; coordinates from the top-left corner in points."""

    WIDTH, HEIGHT = 612, 792

    def __init__(self):
        self.pages: list[dict] = []
        self.new_page()

    def new_page(self):
        self.pages.append({"ops": [], "image": None})

    def text(self, x, y, value, size=10, bold=False, right=False):
        if right:
            x -= _text_width(value, size)
        font = b"/F2" if bold else b"/F1"
        self.pages[-1]["ops"].append(
            b"BT " + font + b" %.1f Tf %.1f %.1f Td " % (size, x, self.HEIGHT - y) + _pdf_string(value) + b" Tj ET")

    def rule(self, x1, y, x2, width=0.6):
        self.pages[-1]["ops"].append(b"%.1f w %.1f %.1f m %.1f %.1f l S" % (width, x1, self.HEIGHT - y, x2, self.HEIGHT - y))

    def shade(self, x, y, w, h, gray=0.9):
        self.pages[-1]["ops"].append(b"%.2f g %.1f %.1f %.1f %.1f re f 0 g" % (gray, x, self.HEIGHT - y - h, w, h))

    def image_page(self, jpeg: bytes, pixel_width: int, pixel_height: int):
        """Replace the current page with a full-page scanned image (no text layer)."""
        self.pages[-1] = {"ops": [b"q %d 0 0 %d 0 0 cm /Im1 Do Q" % (self.WIDTH, self.HEIGHT)],
                          "image": (jpeg, pixel_width, pixel_height)}

    def save(self, path: Path):
        objects: list[bytes] = []

        def add(body: bytes) -> int:
            objects.append(body)
            return len(objects)

        catalog = add(b"")
        pages_id = add(b"")
        f1 = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
        f2 = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
        kids = []
        for page in self.pages:
            content = zlib.compress(b"\n".join(page["ops"]))
            content_id = add(b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(content) + content + b"\nendstream")
            resources = b"/Font << /F1 %d 0 R /F2 %d 0 R >>" % (f1, f2)
            if page["image"]:
                jpeg, w, h = page["image"]
                image_id = add(b"<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace /DeviceRGB "
                               b"/BitsPerComponent 8 /Filter /DCTDecode /Length %d >>\nstream\n" % (w, h, len(jpeg))
                               + jpeg + b"\nendstream")
                resources += b" /XObject << /Im1 %d 0 R >>" % image_id
            kids.append(add(b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 %d %d] /Resources << %s >> /Contents %d 0 R >>"
                            % (pages_id, self.WIDTH, self.HEIGHT, resources, content_id)))
        objects[catalog - 1] = b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id
        objects[pages_id - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(b"%d 0 R" % k for k in kids), len(kids))

        out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        offsets = []
        for number, body in enumerate(objects, start=1):
            offsets.append(len(out))
            out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
        xref = len(out)
        out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
        out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
        out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, catalog, xref)
        path.write_bytes(bytes(out))


ADVICE_LABELS = {
    "en": {"title": "REMITTANCE ADVICE", "payee": "Payee", "columns": ("Invoice / PRO", "Invoice Date", "Invoice Amount", "Discount", "Amount Paid"),
           "total": "Total Payment", "page": "Page", "date_format": "%m/%d/%Y"},
    "es": {"title": "AVISO DE PAGO", "payee": "Beneficiario", "columns": ("Factura / PRO", "Fecha", "Importe", "Descuento", "Importe pagado"),
           "total": "Total pagado", "page": "Página", "date_format": "%d/%m/%Y"},
}


def advice_pdf(path: Path, payer: str, address: list[str], header: list[tuple[str, str]], lines: list[Line],
               stated_total: Decimal | None = None, lang: str = "en", footer: str | None = None) -> None:
    """A text-based remittance advice: payer block, payment details, line table, total, page numbers."""
    labels = ADVICE_LABELS[lang]
    pdf = Pdf()
    x_cols = (50, 190, 365, 450, 560)  # first two left-aligned, the amounts right-aligned at their x
    per_page = 34
    chunks = [lines[i:i + per_page] for i in range(0, len(lines), per_page)] or [[]]
    for page_number, chunk in enumerate(chunks, start=1):
        if page_number > 1:
            pdf.new_page()
        pdf.text(50, 60, payer, 15, bold=True)
        for i, row in enumerate(address):
            pdf.text(50, 78 + i * 13, row, 9)
        pdf.text(560, 60, labels["title"], 14, bold=True, right=True)
        y = 128
        for i, (label, value) in enumerate([(labels["payee"], PAYEE)] + header):
            pdf.text(330, y + i * 14, f"{label}:", 9, bold=True)
            pdf.text(430, y + i * 14, value, 9)
        y = max(y + (len(header) + 1) * 14, 140) + 26
        pdf.shade(45, y - 12, 520, 18)
        for i, column in enumerate(labels["columns"]):
            pdf.text(x_cols[i], y, column, 9, bold=True, right=i >= 2)
        y += 20
        for line in chunk:
            pdf.text(x_cols[0], y, line.pro, 9)
            pdf.text(x_cols[1], y, line.invoice_date.strftime(labels["date_format"]), 9)
            pdf.text(x_cols[2], y, money(line.gross), 9, right=True)
            pdf.text(x_cols[3], y, money(line.discount), 9, right=True)
            pdf.text(x_cols[4], y, money(line.amount), 9, right=True)
            if line.note:
                pdf.text(x_cols[0] + 62, y, line.note, 7)
            y += 15
        if page_number == len(chunks):
            pdf.rule(360, y - 4, 565)
            pdf.text(450, y + 10, labels["total"], 10, bold=True, right=True)
            pdf.text(560, y + 10, money(stated_total if stated_total is not None else total(lines), "$"), 10, bold=True, right=True)
            if footer:
                pdf.text(50, y + 40, footer, 8)
        pdf.text(560, 760, f"{labels['page']} {page_number} / {len(chunks)}", 8, right=True)
    pdf.save(path)


def transfer_confirmation_pdf(path: Path, bank: str, payer: str, amount: Decimal, when: date, reference: str, account_tail: str) -> None:
    """A bank's transfer receipt: proves a payment happened, lists no invoices."""
    pdf = Pdf()
    pdf.shade(0, 0, 612, 70, 0.85)
    pdf.text(50, 42, bank, 18, bold=True)
    pdf.text(50, 110, "ACH Credit Transfer - Confirmation", 14, bold=True)
    rows = [("Status", "Completed"), ("Originator", payer), ("Beneficiary", PAYEE),
            ("Beneficiary account", f"******{account_tail}"), ("Effective date", when.strftime("%B %d, %Y")),
            ("Amount", money(amount, "$") + " USD"), ("Trace / reference number", reference), ("SEC code", "CCD")]
    for i, (label, value) in enumerate(rows):
        pdf.text(50, 150 + i * 20, label, 10, bold=True)
        pdf.text(230, 150 + i * 20, value, 10)
    pdf.text(50, 340, "This confirmation is provided for your records. Please allow 1-2 business days for the funds to post.", 8)
    pdf.save(path)


# --- spreadsheets --------------------------------------------------------------------------------


def remittance_xlsx(path: Path, payer: str, header: list[tuple[str, str]], columns: list[str], rows: list[list],
                    text_columns: tuple[int, ...] = (0,), total_row: list | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Remittance"
    sheet.append([payer])
    sheet["A1"].font = Font(bold=True, size=13)
    for label, value in header:
        sheet.append([label, value])
    sheet.append([])
    sheet.append(columns)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for row in rows:
        sheet.append(row)
        for index in text_columns:
            cell = sheet.cell(row=sheet.max_row, column=index + 1)
            cell.number_format = "@"
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
    if total_row:
        sheet.append(total_row)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
    for column, width in zip("ABCDEF", (18, 14, 16, 12, 16, 24)):
        sheet.column_dimensions[column].width = width
    book.save(path)


# --- images (rendered by render_images.ps1) --------------------------------------------------------


def table_image_job(path: Path, fmt: str, title: str, subtitle: list[str], lines: list[Line], stated_total: Decimal,
                    scan: bool, width: int = 1275, height: int = 1650) -> dict:
    """A remittance table page as an image. 1275 x 1650 px = letter at 150 dpi."""
    items = [{"kind": "rect", "x": 0, "y": 0, "w": width, "h": 120, "color": "#1F3864"},
             {"kind": "text", "text": title, "x": 60, "y": 35, "size": 44, "bold": True, "color": "#FFFFFF"}]
    y = 160
    for row in subtitle:
        items.append({"kind": "text", "text": row, "x": 60, "y": y, "size": 24})
        y += 36
    y += 30
    items.append({"kind": "rect", "x": 50, "y": y - 6, "w": width - 100, "h": 44, "color": "#D9D9D9"})
    header = f"{'PRO / Invoice':<18}{'Invoice Date':<16}{'Amount Paid':>16}"
    items.append({"kind": "text", "text": header, "x": 70, "y": y, "size": 28, "bold": True, "font": "Consolas"})
    y += 56
    for line in lines:
        row = f"{line.pro:<18}{line.invoice_date.strftime('%m/%d/%Y'):<16}{money(line.amount):>16}"
        items.append({"kind": "text", "text": row, "x": 70, "y": y, "size": 28, "font": "Consolas"})
        y += 42
    y += 10
    items.append({"kind": "rect", "x": 600, "y": y, "w": 330, "h": 3, "color": "#000000"})
    items.append({"kind": "text", "text": f"{'Total paid':<18}{'':<16}{money(stated_total, '$'):>16}", "x": 70, "y": y + 14,
                  "size": 28, "bold": True, "font": "Consolas"})
    return {"path": str(path), "format": fmt, "width": width, "height": height, "scan": scan, "items": items}


def logo_job(path: Path, name: str, color: str) -> dict:
    return {"path": str(path), "format": "png", "width": 260, "height": 64, "scan": False,
            "items": [{"kind": "rect", "x": 0, "y": 0, "w": 64, "h": 64, "color": color},
                      {"kind": "text", "text": name[:1], "x": 16, "y": 6, "size": 44, "bold": True, "color": "#FFFFFF"},
                      {"kind": "text", "text": name, "x": 76, "y": 16, "size": 26, "bold": True, "color": color}]}


# --- email bodies --------------------------------------------------------------------------------


def text_table(lines: list[Line], stated_total: Decimal | None = None, separator: str = "  ") -> str:
    rows = [f"{'PRO #':<12}{separator}{'Inv Date':<10}{separator}{'Amount Paid':>12}"]
    rows += [f"{line.pro:<12}{separator}{line.invoice_date.strftime('%m/%d/%Y'):<10}{separator}{money(line.amount):>12}" for line in lines]
    rows.append(f"{'TOTAL':<12}{separator}{'':<10}{separator}{money(stated_total if stated_total is not None else total(lines)):>12}")
    return "\r\n".join(rows)


def signature(name: str, title: str, company: str, phone: str) -> str:
    return f"\r\n\r\nThank you,\r\n\r\n{name}\r\n{title}\r\n{company}\r\n{phone}"


def html(paragraphs: list[str], table: str = "", logo_cid: str | None = None, sign: list[str] | None = None) -> str:
    style = "font-family:Calibri,Arial,sans-serif;font-size:11pt"
    parts = [f'<html><body style="{style}">']
    parts += [f"<p>{p}</p>" for p in paragraphs]
    parts.append(table)
    if sign:
        parts.append("<p>" + "<br>".join(sign) + "</p>")
    if logo_cid:
        parts.append(f'<p><img src="cid:{logo_cid}" width="260" height="64" alt="logo"></p>')
    parts.append("</body></html>")
    return "".join(parts)


def html_table(lines: list[Line]) -> str:
    cell = 'style="border:1px solid #999;padding:3px 8px"'
    right = 'style="border:1px solid #999;padding:3px 8px;text-align:right"'
    rows = [f"<tr><th {cell}>Invoice #</th><th {cell}>Invoice Date</th><th {right}>Invoice Amt</th><th {right}>Paid Amt</th></tr>"]
    rows += [f"<tr><td {cell}>{l.pro}</td><td {cell}>{l.invoice_date.strftime('%m/%d/%Y')}</td>"
             f"<td {right}>{money(l.gross)}</td><td {right}>{money(l.amount)}</td></tr>" for l in lines]
    rows.append(f"<tr><td {cell} colspan=\"3\"><b>Total</b></td><td {right}><b>{money(total(lines))}</b></td></tr>")
    return '<table style="border-collapse:collapse;font-size:10pt">' + "".join(rows) + "</table>"


# --- scenarios -----------------------------------------------------------------------------------


@dataclass
class Message:
    subject: str
    from_name: str
    from_email: str
    sent: datetime
    body_text: str | None = None
    body_html: str | None = None
    attachments: list[dict] = field(default_factory=list)  # {"path", "cid"?}
    embedded: list["Message"] = field(default_factory=list)


@dataclass
class Scenario:
    id: str
    summary: str
    notes: str
    message: Message
    lines: list[Line]
    status: str
    metadata: dict
    issues: list[str]


class Builder:
    def __init__(self, work: Path):
        self.work = work
        self.images: list[dict] = []
        self.scans: list[tuple[Path, Path]] = []  # (jpeg, pdf) pairs wrapped after rendering

    def folder(self, scenario_id: str) -> Path:
        path = self.work / scenario_id
        path.mkdir(parents=True, exist_ok=True)
        return path

    def logo(self, folder: Path, name: str, color: str) -> dict:
        path = folder / "image001.png"
        self.images.append(logo_job(path, name, color))
        return {"path": str(path), "cid": "image001.png@01DB0000.00000000"}


def iso(day: date) -> str:
    return day.isoformat()


def scenarios(b: Builder) -> list[Scenario]:
    out: list[Scenario] = []

    # 01 - plain-text table in the email body
    d = Data(101)
    paid = date(2026, 10, 5)
    lines = d.lines(6, paid)
    ref = f"ACH{d.digits(10)}"
    out.append(Scenario(
        "01-body-text-table", "Remittance as a plain-text table in the email body",
        "Body-only remittance: no attachments. Fixed-width text table with a TOTAL row.",
        Message(f"Payment Remittance - {paid:%m/%d/%Y}", "Fabrikam Accounts Payable", "ap@fabrikam.example", datetime(2026, 10, 5, 14, 12),
                body_text=(f"Hello,\r\n\r\nPlease see below for the ACH payment sent {paid:%m/%d/%Y}, reference {ref}.\r\n\r\n"
                           + text_table(lines) + signature("Maria Lopez", "Accounts Payable Specialist", "Fabrikam, Inc.", "(555) 010-2231"))),
        lines, "Completed", {"payerName": "Fabrikam, Inc.", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref}, []))

    # 02 - HTML table in the body, with a signature logo (inline image to ignore)
    d = Data(102)
    paid = date(2026, 10, 6)
    lines = d.lines(9, paid, discount=True)
    folder = b.folder("02-body-html-table")
    logo = b.logo(folder, "Northwind", "#2E7D32")
    ref = f"PMT-{d.digits(7)}"
    out.append(Scenario(
        "02-body-html-table", "Remittance as an HTML table in the body, with a signature logo",
        "HTML-formatted body table with invoice and paid amounts (some lines carry a 2% discount). The signature logo is an inline image that is not a remittance source.",
        Message(f"Northwind Traders - ACH Payment {ref}", "Northwind Traders AP", "payables@northwind.example", datetime(2026, 10, 6, 15, 40),
                body_html=html([f"Good afternoon,", f"An ACH payment of <b>{money(total(lines), '$')}</b> was released on {paid:%B %d, %Y} (payment {ref}). Details:"],
                               html_table(lines), logo["cid"], ["Kind regards,", "Daniel Park", "Accounts Payable", "Northwind Traders"]),
                attachments=[logo]),
        lines, "Completed", {"payerName": "Northwind Traders", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref}, []))

    # 03 - text-based PDF remittance advice
    d = Data(103)
    paid = date(2026, 10, 7)
    lines = d.lines(14, paid, discount=True)
    folder = b.folder("03-pdf-remittance-advice")
    ref = f"{d.digits(6)}"
    pdf = folder / f"Remittance_Advice_{ref}.pdf"
    advice_pdf(pdf, "Contoso Manufacturing", ["1200 Industrial Parkway", "Springfield, ST 00001"],
               [("Payment date", f"{paid:%m/%d/%Y}"), ("Payment method", "ACH"), ("Payment number", ref), ("Vendor number", "V-20417")], lines)
    out.append(Scenario(
        "03-pdf-remittance-advice", "Short email with a one-page PDF remittance advice",
        "Typical ERP remittance advice PDF (text layer) with gross, discount and paid columns.",
        Message(f"Remittance Advice {ref}", "Contoso Manufacturing", "noreply-payments@contoso.example", datetime(2026, 10, 7, 11, 5),
                body_text="This is an automated remittance advice. Please see the attached document for payment details.\r\n\r\nDo not reply to this message.",
                attachments=[{"path": str(pdf)}]),
        lines, "Completed", {"payerName": "Contoso Manufacturing", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref}, []))

    # 04 - scanned PDF (image only, needs OCR)
    d = Data(104)
    paid = date(2026, 10, 8)
    lines = d.lines(10, paid)
    folder = b.folder("04-scanned-pdf")
    ref = f"CHK{d.digits(6)}"
    jpeg = folder / "scan.jpg"
    b.images.append(table_image_job(jpeg, "jpeg", "Tailspin Toys - Payment Notice",
                                    [f"Payment date: {paid:%m/%d/%Y}", f"Payment ID: {ref}", f"Paid to: {PAYEE}"], lines, total(lines), scan=True))
    b.scans.append((jpeg, folder / "Scan_20261008.pdf"))
    out.append(Scenario(
        "04-scanned-pdf", "Scanned PDF remittance (image only, no text layer)",
        "The attachment is a tilted, speckled page image inside a PDF; extraction depends on OCR.",
        Message("Scanned document", "Tailspin Toys Accounting", "accounting@tailspintoys.example", datetime(2026, 10, 8, 9, 30),
                body_text="Please find the remittance for our payment attached.\r\n\r\nTailspin Toys Accounting",
                attachments=[{"path": str(folder / "Scan_20261008.pdf")}]),
        lines, "Completed", {"payerName": "Tailspin Toys", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref}, []))

    # 05 - PNG screenshot from a payment portal
    d = Data(105)
    paid = date(2026, 10, 9)
    lines = d.lines(7, paid)
    folder = b.folder("05-image-png")
    ref = f"{d.digits(9)}"
    png = folder / "payment_details.png"
    b.images.append(table_image_job(png, "png", "Payment Details", [f"Payer: Wingtip Toys", f"Settlement date: {paid:%m/%d/%Y}",
                                                                    f"Confirmation #: {ref}"], lines, total(lines), scan=False))
    out.append(Scenario(
        "05-image-png", "Remittance as a PNG screenshot of a payment portal",
        "A clean image attachment (no PDF, no text) - extraction depends on OCR.",
        Message("Payment confirmation", "Wingtip Toys Finance", "finance@wingtiptoys.example", datetime(2026, 10, 9, 16, 2),
                body_text="Hi team,\r\n\r\nScreenshot of today's payment from our bank portal attached.\r\n\r\nThanks,\r\nPriya",
                attachments=[{"path": str(png)}]),
        lines, "Completed", {"payerName": "Wingtip Toys", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref}, []))

    # 06 - TIFF with no file extension
    d = Data(106)
    paid = date(2026, 10, 12)
    lines = d.lines(8, paid)
    folder = b.folder("06-tiff-no-extension")
    ref = f"{d.digits(8)}"
    tiff = folder / f"REMIT_ADVICE_{ref}"
    b.images.append(table_image_job(tiff, "tiff", "Adventure Works - Remittance", [f"Date: {paid:%m/%d/%Y}", f"Remittance no.: {ref}"],
                                    lines, total(lines), scan=True))
    out.append(Scenario(
        "06-tiff-no-extension", "Faxed TIFF remittance whose attachment name has no extension",
        "The file type must be recognized from its content (TIFF signature), then read with OCR.",
        Message(f"Remittance {ref}", "Adventure Works", "remittance@adventure-works.example", datetime(2026, 10, 12, 8, 47),
                body_text="Remittance attached (open with an image viewer).",
                attachments=[{"path": str(tiff)}]),
        lines, "Completed", {"payerName": "Adventure Works", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref}, []))

    # 07 - Excel remittance
    d = Data(107)
    paid = date(2026, 10, 13)
    lines = d.lines(40, paid, discount=True)
    folder = b.folder("07-excel-remittance")
    ref = f"EFT{d.digits(7)}"
    xlsx = folder / f"Payment_{ref}.xlsx"
    remittance_xlsx(xlsx, "Litware, Inc. - Payment Detail", [("Payment date", f"{paid:%m/%d/%Y}"), ("Payment ref", ref), ("Payee", PAYEE)],
                    ["Invoice Number", "Invoice Date", "Gross Amount", "Discount", "Net Paid"],
                    [[l.pro, l.invoice_date.strftime("%m/%d/%Y"), float(l.gross), float(l.discount), float(l.amount)] for l in lines],
                    total_row=["Total", "", "", "", float(total(lines))])
    out.append(Scenario(
        "07-excel-remittance", "Excel remittance with a header block above the table",
        "40-line .xlsx: payer and payment details above the table, then gross / discount / net columns and a total row.",
        Message(f"Litware Payment {ref}", "Litware AP", "ap@litware.example", datetime(2026, 10, 13, 13, 20),
                body_text=f"Attached is the detail for payment {ref}, sent via EFT on {paid:%m/%d/%Y}." + signature("Chris Novak", "AP Analyst", "Litware, Inc.", "(555) 010-8890"),
                attachments=[{"path": str(xlsx)}]),
        lines, "Completed", {"payerName": "Litware, Inc.", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref}, []))

    # 08 - leading zeros in PROs (Excel text cells)
    d = Data(108)
    paid = date(2026, 10, 14)
    lines = d.lines(12, paid, leading_zero=True)
    folder = b.folder("08-leading-zeros")
    ref = f"{d.digits(10)}"
    xlsx = folder / "remit detail.xlsx"
    remittance_xlsx(xlsx, "Proseware Logistics", [("Check/EFT #", ref), ("Pay date", f"{paid:%m/%d/%Y}")],
                    ["PRO Number", "Ship Date", "Amount"],
                    [[l.pro, l.invoice_date.strftime("%m/%d/%Y"), float(l.amount)] for l in lines])
    out.append(Scenario(
        "08-leading-zeros", "PRO numbers with leading zeros, stored as text in Excel",
        "Every PRO starts with 00; the leading zeros must be kept. No total row - the total is stated only in the body.",
        Message("ACH remittance", "Proseware Logistics", "payments@proseware.example", datetime(2026, 10, 14, 10, 15),
                body_text=f"Payment of {money(total(lines), '$')} sent today, EFT # {ref}. Detail attached.",
                attachments=[{"path": str(xlsx)}]),
        lines, "Completed", {"payerName": "Proseware Logistics", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref}, []))

    # 09 - very large table in the body
    d = Data(109)
    paid = date(2026, 10, 15)
    lines = d.lines(1200, paid)
    ref = f"WIRE{d.digits(8)}"
    out.append(Scenario(
        "09-large-body-table", "1,200-line tab-separated table in the email body",
        "A broker-scale remittance pasted into the body (tab-separated). Tests large-table handling and the total to the cent.",
        Message(f"Wide World Importers - Payment {ref} - {len(lines)} items", "Wide World Importers Carrier Pay", "carrierpay@wideworldimporters.example",
                datetime(2026, 10, 15, 18, 55),
                body_text=(f"Carrier payment {ref}, paid {paid:%m/%d/%Y}, total {money(total(lines), '$')}.\r\n\r\n" + text_table(lines, separator="\t")
                           + "\r\n\r\nWide World Importers - Carrier Settlement")),
        lines, "Completed", {"payerName": "Wide World Importers", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref}, []))

    # 10 - the same lines in a PDF and in Excel (corroboration, not duplicates)
    d = Data(110)
    paid = date(2026, 10, 16)
    lines = d.lines(11, paid)
    folder = b.folder("10-pdf-and-excel-same-lines")
    ref = f"{d.digits(7)}"
    pdf, xlsx = folder / f"Remittance {ref}.pdf", folder / f"Remittance {ref}.xlsx"
    advice_pdf(pdf, "Coho Winery", ["88 Vineyard Road", "Riverton, ST 00002"], [("Payment date", f"{paid:%m/%d/%Y}"), ("Payment no.", ref)], lines)
    remittance_xlsx(xlsx, "Coho Winery", [("Payment no.", ref)], ["Invoice", "Date", "Amount Paid"],
                    [[l.pro, l.invoice_date.strftime("%m/%d/%Y"), float(l.amount)] for l in lines], total_row=["Total", "", float(total(lines))])
    out.append(Scenario(
        "10-pdf-and-excel-same-lines", "The same remittance as both a PDF and an Excel file",
        "Two sources with identical lines: they corroborate each other and must not be counted twice or flagged as duplicate payments.",
        Message(f"Payment {ref} - remittance (PDF + Excel)", "Coho Winery", "ap@cohowinery.example", datetime(2026, 10, 16, 12, 0),
                body_text="Attached in PDF and Excel format for your convenience.",
                attachments=[{"path": str(pdf)}, {"path": str(xlsx)}]),
        lines, "Completed", {"payerName": "Coho Winery", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref}, []))

    # 11 - wire fee short pay
    d = Data(111)
    paid = date(2026, 10, 19)
    lines = d.lines(5, paid)
    fee = Decimal("15.00")
    received = total(lines) - fee
    folder = b.folder("11-wire-fee-short-pay")
    ref = f"FT{d.digits(10)}"
    pdf = folder / "remittance.pdf"
    advice_pdf(pdf, "Fourth Coffee Ltd.", ["Unit 4, Harbour Estate", "Port Town, ST 00003"],
               [("Value date", f"{paid:%m/%d/%Y}"), ("Wire reference", ref)], lines, stated_total=received,
               footer=f"Invoices total {money(total(lines), '$')} less bank wire fee {money(fee, '$')}.")
    out.append(Scenario(
        "11-wire-fee-short-pay", "International wire: the amount paid is the invoices less a bank fee",
        f"Invoice lines total {money(total(lines))}; the stated payment is {money(received)} (a {money(fee)} wire fee deducted).",
        Message(f"Wire payment {ref}", "Fourth Coffee Treasury", "treasury@fourthcoffee.example", datetime(2026, 10, 19, 7, 25),
                body_text=f"Dear Sir/Madam,\r\n\r\nWe have remitted {money(received, '$')} by wire today. Our bank deducted a {money(fee, '$')} fee.\r\nRemittance attached.\r\n\r\nRegards,\r\nFourth Coffee Treasury",
                attachments=[{"path": str(pdf)}]),
        lines, "ReviewRequired", {"payerName": "Fourth Coffee Ltd.", "remittanceDate": iso(paid), "paymentTotal": float(received), "paymentReference": ref},
        ["payer:payment_total_mismatch"]))

    # 12 - credit memo (negative line)
    d = Data(112)
    paid = date(2026, 10, 20)
    lines = d.lines(6, paid)
    lines.append(Line(d.pro(), Decimal("-125.00"), paid - timedelta(days=20), note="Credit memo - damaged freight claim"))
    folder = b.folder("12-credit-memo-negative")
    ref = f"{d.digits(6)}"
    pdf = folder / f"RA_{ref}.pdf"
    advice_pdf(pdf, "Alpine Ski House", ["400 Summit Drive", "Highpoint, ST 00004"], [("Payment date", f"{paid:%m/%d/%Y}"), ("Payment #", ref)], lines)
    out.append(Scenario(
        "12-credit-memo-negative", "Remittance that nets a credit memo (negative line)",
        "One line is -125.00 (a claim credit taken against the payment); the total is net of the credit.",
        Message(f"Remittance {ref}", "Alpine Ski House AP", "ap@alpineskihouse.example", datetime(2026, 10, 20, 14, 45),
                body_text="Please see attached remittance. Note we have deducted credit memo for the damaged shipment claim.",
                attachments=[{"path": str(pdf)}]),
        lines, "ReviewRequired", {"payerName": "Alpine Ski House", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref},
        ["payer:negative_amount_paid"]))

    # 13 - forwarded thread with two payments (embedded emails + bank confirmations)
    d = Data(113)
    folder = b.folder("13-forwarded-multi-payment")
    inner, all_lines = [], []
    for index, (paid, hour) in enumerate(((date(2026, 10, 14), 10), (date(2026, 10, 16), 15)), start=1):
        lines = d.lines(4 + index, paid)
        all_lines += lines
        ref = f"{d.digits(15)}"
        proof = folder / f"ACH_Confirmation_{index}.pdf"
        transfer_confirmation_pdf(proof, "Woodgrove Bank", "Trey Research", total(lines), paid, ref, d.digits(4))
        inner_msg = Message(f"ACH payment {paid:%m/%d} - {money(total(lines), '$')}", "Trey Research AP", "ap@treyresearch.example",
                            datetime(paid.year, paid.month, paid.day, hour, 10),
                            body_text=f"Payment released {paid:%m/%d/%Y} for the following:\r\n\r\n" + text_table(lines) + "\r\n\r\nBank confirmation attached.",
                            attachments=[{"path": str(proof)}])
        inner.append(inner_msg)
    out.append(Scenario(
        "13-forwarded-multi-payment", "Forwarded thread carrying two separate payments, each with a bank confirmation",
        "The outer email attaches two earlier emails (embedded .msg). Each has its own line table and an ACH transfer confirmation PDF that lists no invoices. One email, two bank credits.",
        Message("FW: ACH payments this week", "Trey Research AP", "ap@treyresearch.example", datetime(2026, 10, 16, 16, 30),
                body_text="Hi,\r\n\r\nForwarding both payments sent this week with their confirmations, as requested.\r\n\r\nThanks,\r\nSam Rivera\r\nTrey Research",
                embedded=inner),
        all_lines, "ReviewRequired", {"payerName": "Trey Research"}, ["payer:multiple_payments_in_email"]))

    # 14 - dispute / collections inquiry thread
    d = Data(114)
    pros = [d.pro(), d.pro(), d.pro()]
    unapplied = d.amount(900, 2500)
    thread = (
        "Hi Jordan,\r\n\r\nWe show both of those invoices as paid in full last month. The "
        f"{money(unapplied, '$')} you see as unapplied was meant for PRO {pros[2]}. Can you apply it there and send us an updated statement?\r\n\r\n"
        "Thanks,\r\nAlex Chen\r\nLucerne Publishing - Accounts Payable\r\n\r\n"
        "From: Jordan Miles <collections@foundryfreight.example>\r\nSent: Monday, October 19, 2026 9:02 AM\r\nTo: Lucerne Publishing AP <ap@lucernepublishing.example>\r\n"
        "Subject: Past due - PRO " + pros[0] + " and " + pros[1] + "\r\n\r\n"
        f"Hello,\r\n\r\nOur records show PRO {pros[0]} and PRO {pros[1]} are past due, and we have an unapplied payment of {money(unapplied, '$')} "
        "on your account. Could you tell us how that payment should be applied?\r\n\r\nRegards,\r\nJordan Miles\r\nCollections, " + PAYEE)
    out.append(Scenario(
        "14-dispute-thread", "Collections inquiry thread - discussion about payments, not a remittance",
        "The payer and the collections team discuss past-due PROs and where to apply an unapplied amount. PROs and amounts appear, but nothing is being remitted.",
        Message(f"RE: Past due - PRO {pros[0]} and {pros[1]}", "Lucerne Publishing AP", "ap@lucernepublishing.example", datetime(2026, 10, 20, 11, 18),
                body_text=thread),
        [], "ReviewRequired", {"payerName": "Lucerne Publishing"}, ["payer:remittance_inquiry_thread"]))

    # 15 - no remittance detail at all
    out.append(Scenario(
        "15-no-remittance", "Payment notification with no remittance detail",
        "The payer says a payment was sent but gives no invoices, PROs or amounts.",
        Message("Payment sent", "Margie's Travel", "office@margiestravel.example", datetime(2026, 10, 21, 17, 3),
                body_text="Hi,\r\n\r\nJust letting you know we sent payment today for the open invoices. Let me know if you need anything else.\r\n\r\nBest,\r\nRenee"),
        [], "ReviewRequired", {"payerName": "Margie's Travel"}, ["payer:no_remittance_in_email"]))

    # 16 - remittance only behind a portal link
    d = Data(116)
    ref = f"{d.digits(10)}"
    out.append(Scenario(
        "16-portal-link", "Remittance detail available only behind a supplier-portal link",
        "The email states the amount and links to a portal for the detail; the link is not fetched (the domain is reserved and never resolves).",
        Message(f"Payment Notification - {ref}", "Relecloud Supplier Portal", "no-reply@relecloud.example", datetime(2026, 10, 21, 6, 0),
                body_html=html(["Dear Supplier,", f"A payment (reference <b>{ref}</b>) has been issued to {PAYEE}.",
                                f'View the full remittance detail in the supplier portal: <a href="https://suppliers.relecloud.example/remittance/{ref}">'
                                f"https://suppliers.relecloud.example/remittance/{ref}</a>",
                                "This mailbox is not monitored."])),
        [], "ReviewRequired", {"paymentReference": ref}, ["payer:external_link_not_fetched"]))

    # 17 - Spanish-language remittance (PDF + body)
    d = Data(117)
    paid = date(2026, 10, 22)
    lines = d.lines(8, paid)
    folder = b.folder("17-spanish-remittance")
    ref = f"SPEI{d.digits(10)}"
    pdf = folder / f"Aviso_de_pago_{ref}.pdf"
    advice_pdf(pdf, "Proveedora Litware S.A. de C.V.", ["Av. Industria 450", "Monterrey, N.L., México"],
               [("Fecha de pago", f"{paid:%d/%m/%Y}"), ("Referencia", ref), ("Moneda", "USD")], lines, lang="es")
    out.append(Scenario(
        "17-spanish-remittance", "Spanish-language email and PDF remittance (day-first dates)",
        "Labels in Spanish (Factura, Importe pagado, Total pagado) and dates written dd/mm/yyyy.",
        Message(f"Aviso de pago {ref}", "Cuentas por Pagar - Proveedora Litware", "cxp@litware-mx.example", datetime(2026, 10, 22, 16, 40),
                body_text=(f"Buen día,\r\n\r\nAnexamos el aviso de pago correspondiente a la transferencia del {paid:%d/%m/%Y} por un total de "
                           f"{money(total(lines), '$')} USD.\r\n\r\nSaludos cordiales,\r\nCuentas por Pagar\r\nProveedora Litware S.A. de C.V."),
                attachments=[{"path": str(pdf)}]),
        lines, "Completed", {"payerName": "Proveedora Litware S.A. de C.V.", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref}, []))

    # 18 - EDI 820 / CCD+ addenda in the body
    d = Data(118)
    paid = date(2026, 10, 22)
    lines = d.lines(5, paid)
    trace = f"{d.digits(15)}"
    segments = [f"BPR*C*{total(lines):.2f}*C*ACH*CCP*01*000000000*DA*0000000000**01*111000025*DA*0000000000*{paid:%Y%m%d}~",
                f"TRN*1*{trace}*1000000000~", "N1*PR*VANARSDEL LTD~", f"N1*PE*{PAYEE.upper()}~"]
    segments += [f"RMR*IV*{l.pro}**{l.amount:.2f}*{l.gross:.2f}~" for l in lines]
    out.append(Scenario(
        "18-edi-820-body", "ACH CCD+/CTX remittance as EDI 820 segments in the body",
        "Bank-generated notice with raw X12 820 segments: BPR carries the total and date, TRN the trace number, one RMR*IV segment per PRO.",
        Message(f"ACH Credit Received - Trace {trace}", "Woodgrove Bank Notifications", "alerts@woodgrovebank.example", datetime(2026, 10, 22, 12, 30),
                body_text=(f"An ACH credit of {money(total(lines), '$')} from VANARSDEL LTD was received on {paid:%m/%d/%Y}.\r\n\r\nAddenda:\r\n"
                           + "\r\n".join(segments))),
        lines, "Completed", {"payerName": "VanArsdel Ltd", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": trace}, []))

    # 19 - the same PRO and amount twice in one source (possible double payment)
    d = Data(119)
    paid = date(2026, 10, 23)
    lines = d.lines(6, paid)
    lines.insert(4, Line(lines[1].pro, lines[1].amount, lines[1].invoice_date))
    ref = f"{d.digits(8)}"
    out.append(Scenario(
        "19-duplicate-line", "A PRO paid twice for the same amount in one remittance",
        "One PRO and amount appear on two lines of the same table; the payment total includes both.",
        Message(f"ACH Remittance {ref}", "A. Datum Corporation", "ap@adatum.example", datetime(2026, 10, 23, 9, 50),
                body_text=f"Payment {ref} sent {paid:%m/%d/%Y}:\r\n\r\n" + text_table(lines) + signature("Lee Wong", "AP", "A. Datum Corporation", "(555) 010-4410")),
        lines, "ReviewRequired", {"payerName": "A. Datum Corporation", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref},
        ["payer:duplicate_payment"]))

    # 20 - the same PRO with two different amounts
    d = Data(120)
    paid = date(2026, 10, 23)
    lines = d.lines(7, paid)
    lines.append(Line(lines[2].pro, (lines[2].amount / 10).quantize(CENT), lines[2].invoice_date))
    folder = b.folder("20-conflicting-amounts")
    ref = f"{d.digits(9)}"
    xlsx = folder / f"Remit_{ref}.xlsx"
    remittance_xlsx(xlsx, "Trey Research Labs", [("Payment ID", ref), ("Payment Date", f"{paid:%m/%d/%Y}")], ["Invoice", "Invoice Date", "Paid"],
                    [[l.pro, l.invoice_date.strftime("%m/%d/%Y"), float(l.amount)] for l in lines], total_row=["Total", "", float(total(lines))])
    out.append(Scenario(
        "20-conflicting-amounts", "One PRO on two lines with different amounts",
        "A PRO is listed twice with different paid amounts (for example a keying error or a partial payment); the total includes both.",
        Message(f"Remittance - Payment ID {ref}", "Trey Research Labs AP", "ap-labs@treyresearch.example", datetime(2026, 10, 23, 14, 5),
                body_text="Remittance detail attached.", attachments=[{"path": str(xlsx)}]),
        lines, "ReviewRequired", {"payerName": "Trey Research Labs", "remittanceDate": iso(paid), "paymentTotal": float(total(lines)), "paymentReference": ref},
        ["payer:conflicting_payment_line"]))

    return out


# --- build ---------------------------------------------------------------------------------------


def run_powershell(script: Path, json_path: Path) -> None:
    subprocess.run(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), str(json_path)], check=True)


def manifest_entries(message: Message, path: Path, work: Path, prefix: str) -> list[dict]:
    """Inner (embedded) messages first, so the outer message can attach their .msg files."""
    entries, attachments = [], list(message.attachments)
    for index, inner in enumerate(message.embedded, start=1):
        inner_path = work / f"{prefix}-embedded-{index}.msg"
        entries += manifest_entries(inner, inner_path, work, f"{prefix}-{index}")
        attachments.append({"path": str(inner_path)})
    entries.append({
        "path": str(path), "subject": message.subject, "to": PAYEE_EMAIL,
        "fromName": message.from_name, "fromEmail": message.from_email, "sentUtc": message.sent.strftime("%Y-%m-%d %H:%M"),
        "bodyText": message.body_text, "bodyHtml": message.body_html, "attachments": attachments,
    })
    return entries


def write_labels(scenario: Scenario, folder: Path) -> None:
    label = {
        "schemaVersion": "1.0", "id": scenario.id, "source": f"{scenario.id}.msg", "verified": None,
        "notes": f"Synthetic sample. {scenario.summary}. {scenario.notes}",
        "expected": {"status": scenario.status, "metadata": scenario.metadata, "issues": scenario.issues},
    }
    (folder / "expected.json").write_text(json.dumps(label, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    with (folder / "expected-lines.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["pro", "amount"])
        for line in scenario.lines:
            writer.writerow([line.pro, f"{line.amount:.2f}"])


def write_catalog(out: Path, built: list[Scenario]) -> None:
    rows = ["# Synthetic remittance samples", "",
            "Generated by `scripts/synthetic_samples/generate.py`. Every company, person, PRO and amount is invented; "
            f"the payee is the fictional **{PAYEE}** (add it to `RECEIVING_ORG_ALIASES` when testing a configured deployment).", "",
            "Each folder holds the `.msg`, `expected.json` and `expected-lines.csv` in the golden-set layout. Labels are unverified: "
            "the lines, totals and dates are exact by construction; the expected status and issues are the intended behavior.", "",
            "| Sample | Scenario | Lines | Expected |", "|---|---|---|---|"]
    for s in built:
        rows.append(f"| `{s.id}` | {s.summary} | {len(s.lines)} | {s.status}{' - ' + ', '.join(s.issues) if s.issues else ''} |")
    rows += ["", "Regenerate with `.venv\\Scripts\\python.exe scripts/synthetic_samples/generate.py` (Windows with Outlook).", ""]
    (out / "README.md").write_text("\n".join(rows), encoding="utf-8")


def local_identity() -> list[str]:
    """The Outlook profile's user name and address, and the Windows user name: Outlook stamps the
    profile identity into every saved item (last modifier, Internet account name and stamp)."""
    import os

    command = ("$u = (New-Object -ComObject Outlook.Application).Session.CurrentUser; $u.Name; "
               "try { $u.AddressEntry.GetExchangeUser().PrimarySmtpAddress } catch { }; "
               "try { $u.AddressEntry.Address } catch { }")
    result = subprocess.run(["powershell.exe", "-NoProfile", "-Command", command], capture_output=True, text=True)
    values = [v.strip() for v in result.stdout.splitlines() if v.strip()] + [os.environ.get("USERNAME", "")]
    tokens = set()
    for value in values:
        tokens.add(value)
        tokens.update(part for part in value.replace("@", " ").replace(".", " ").replace("/", " ").replace("=", " ").split() if len(part) >= 5)
    common = {"exchange", "administrative", "group", "recipients", "exchangelabs", "fydibohf23spdlt"}
    return sorted((t for t in tokens if len(t) >= 4 and t.lower() not in common), key=len, reverse=True)


IDENTITY_STREAMS = ("__substg1.0_3FF8001F", "__substg1.0_3FFA001F")  # creator, last modifier


def scrub(path: Path, tokens: list[str]) -> int:
    """Overwrite, in place, every string stream that carries the local identity (at any depth,
    including embedded messages) with a neutral value of the same length. Returns streams changed."""
    import olefile

    lowered = [t.lower() for t in tokens]
    changed = 0
    with olefile.OleFileIO(str(path), write_mode=True) as ole:
        for entry in ole.listdir(streams=True, storages=False):
            name = entry[-1]
            if not name.startswith("__substg1.0_") or not name.endswith(("001F", "001E")):
                continue
            data = ole.openstream(entry).read()
            text = data.decode("utf-16-le" if name.endswith("001F") else "latin-1", errors="ignore").lower()
            if name not in IDENTITY_STREAMS and not any(t in text for t in lowered):
                continue
            if name.endswith("001F"):
                filler = ("synthetic-sample" + " " * len(data))[: len(data) // 2].encode("utf-16-le")
                filler += b"\x00" * (len(data) - len(filler))
            else:
                filler = ("synthetic-sample" + " " * len(data))[: len(data)].encode("latin-1")
            ole.write_stream(entry, filler)
            changed += 1
    return changed


def verify(out: Path, built: list[Scenario], forbidden: list[str]) -> list[str]:
    """Parse every .msg with the pipeline's own parser and scan it for local-user identity strings."""
    sys.path.insert(0, str(REPO / "apps" / "remittance-processor"))
    from msg_parser.parser import parse_msg

    problems = []
    for s in built:
        path = out / s.id / f"{s.id}.msg"
        content = path.read_bytes()
        for word in forbidden:
            if word and (word.lower().encode("utf-16-le") in content.lower() or word.lower().encode() in content.lower()):
                problems.append(f"{s.id}: contains local identity string '{word}'")
        parsed = parse_msg(content, path.name)
        kinds = ", ".join(f"{a.name}{' (inline)' if a.is_inline else ''}{' in ' + a.embedded_in if a.embedded_in else ''}" for a in parsed.attachments) or "none"
        print(f"  {s.id:30} from={parsed.sender_email} date={(parsed.sent_date or '')[:16]} body={len(parsed.body or '')} chars; attachments: {kinds}")
        if not parsed.sender_email or not parsed.sent_date:
            problems.append(f"{s.id}: sender or date missing")
    return problems


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--out", default=str(REPO / "samples" / "synthetic"))
    parser.add_argument("--only", nargs="*")
    args = parser.parse_args()

    out = Path(args.out).resolve()
    work = out / "_build"
    if work.exists():
        shutil.rmtree(work)
    work.mkdir(parents=True)
    builder = Builder(work)
    built = [s for s in scenarios(builder) if not args.only or s.id in args.only]

    # 1. images (PNG / TIFF / JPEG) - then wrap the scanned JPEGs into image-only PDFs
    jobs = work / "images.json"
    jobs.write_text(json.dumps(builder.images), encoding="utf-8")
    run_powershell(HERE / "render_images.ps1", jobs)
    for jpeg, pdf_path in builder.scans:
        pdf = Pdf()
        pdf.image_page(jpeg.read_bytes(), 1275, 1650)
        pdf.save(pdf_path)

    # 2. messages, through Outlook
    entries = []
    for s in built:
        folder = out / s.id
        if folder.exists():
            shutil.rmtree(folder)
        folder.mkdir(parents=True)
        entries += manifest_entries(s.message, folder / f"{s.id}.msg", work, s.id)
        write_labels(s, folder)
    manifest = work / "messages.json"
    manifest.write_text(json.dumps(entries, ensure_ascii=False), encoding="utf-8")
    run_powershell(HERE / "build_msgs.ps1", manifest)
    write_catalog(out, built)

    # 3. remove the local Outlook identity, then check every file with the pipeline's parser
    tokens = local_identity()
    scrubbed = sum(scrub(out / s.id / f"{s.id}.msg", tokens) for s in built)
    print(f"Built {len(built)} sample(s) in {out}; scrubbed {scrubbed} identity stream(s)")
    problems = verify(out, built, tokens)
    for problem in problems:
        print("PROBLEM:", problem)
    if not problems:
        for scratch in [jobs, manifest, *work.glob("*.msg")]:
            scratch.unlink()
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
