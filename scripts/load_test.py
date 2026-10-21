"""Synthetic ACH remittance load generator.

Generates unique, neutral (fictional) remittance emails as MIME ``.eml`` files - with and without
attachments - and submits them through the real intake path: UI/API -> Blob `incoming/` ->
Service Bus -> Durable orchestration. Use it to demo a day's volume
(e.g. 1,000 messages compressed into an hour) or a morning burst, and watch the Live Pipeline page drain the queue.

Every message is unique (content-hash processingId), so nothing is collapsed as a duplicate unless
you ask for duplicates with --duplicate-rate.

    Mix (default)  body-only 30% | PDF 35% | CSV 15% | XLSX 10% | exceptions 10%
    Exceptions     total mismatch (short pay), portal link (not fetched), conflicting amounts

BILLABLE: every PDF/body message calls Content Understanding (gpt-5.5) and every exception calls
the assurance agent. Runs above 50 messages require --yes.

Examples:
    # 25 messages as fast as the API accepts them
    python scripts/load_test.py --api https://<ui-fqdn> --count 25

    # Simulated day compressed into one hour (1,000 msgs spread evenly)
    python scripts/load_test.py --api https://<ui-fqdn> --count 1000 --duration-minutes 60 --yes

    # Morning burst: 400 messages, 8 concurrent senders, no pacing
    python scripts/load_test.py --api https://<ui-fqdn> --count 400 --concurrency 8 --yes

    # Just write the files (no submission) to inspect them
    python scripts/load_test.py --count 10 --out-dir ./samples/synthetic --dry-run
"""

from __future__ import annotations

import argparse
import io
import json
import random
import sys
import threading
import time
import urllib.error
import urllib.request
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import format_datetime, make_msgid
from pathlib import Path

PAYERS = [
    "Acme Transport Co", "Northwind Distribution", "Contoso Retail Group", "Fabrikam Industrial",
    "Tailspin Outdoor Supply", "Wide World Importers", "Litware Building Products", "Adventure Works Cycles",
    "Proseware Packaging", "Woodgrove Grocers", "Lamna Health Supply", "Alpine Ski House",
    "Blue Yonder Appliances", "Coho Winery", "Margie's Travel Goods", "Trey Research Labs",
]

MIX = {"body": 30, "pdf": 35, "csv": 15, "xlsx": 10, "exception": 10}


# --- remittance data -----------------------------------------------------------------------------


def make_remittance(rng: random.Random) -> dict:
    payer = rng.choice(PAYERS)
    line_count = rng.choices([1, 2, 3, 5, 8, 12, 25, 40], weights=[10, 14, 16, 18, 16, 12, 9, 5])[0]
    lines = []
    for _ in range(line_count):
        pro = f"{rng.randint(0, 999_999_999):09d}"  # leading zeros preserved as text
        lines.append({"pro": pro, "amount": round(rng.uniform(85, 4200), 2)})
    remit_date = date.today() - timedelta(days=rng.randint(0, 6))
    return {
        "payer": payer,
        "date": remit_date.isoformat(),
        "reference": f"ACH{rng.randint(10_000_000, 99_999_999)}",
        "lines": lines,
        "total": round(sum(l["amount"] for l in lines), 2),
        "nonce": uuid.uuid4().hex[:12],
    }


def remittance_text(r: dict, *, total=None) -> str:
    rows = "\n".join(f"  {l['pro']:<12} {l['amount']:>12,.2f}" for l in r["lines"])
    return (
        f"Payer: {r['payer']}\n"
        f"Payment date: {r['date']}\n"
        f"ACH trace / reference: {r['reference']}\n"
        f"Payment total: USD {total if total is not None else r['total']:,.2f}\n\n"
        f"  PRO number   Amount paid\n{rows}\n"
    )


# --- attachment builders ---------------------------------------------------------------------------


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(lines_of_text: list[str]) -> bytes:
    """Minimal single-font text PDF (no dependencies). Long remittances flow onto more pages."""
    per_page = 48
    pages = [lines_of_text[i : i + per_page] for i in range(0, len(lines_of_text), per_page)] or [[""]]
    objects: list[bytes] = []
    font_obj = 3
    page_ids = []
    content_ids = []
    next_id = 4
    for _ in pages:
        page_ids.append(next_id)
        content_ids.append(next_id + 1)
        next_id += 2
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>")
    for page_lines, page_id, content_id in zip(pages, page_ids, content_ids):
        text_ops = "\n".join(f"({_pdf_escape(line)}) Tj T*" for line in page_lines)
        stream = f"BT /F1 10 Tf 14 TL 50 760 Td\n{text_ops}\nET".encode("latin-1", "replace")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_obj} 0 R >> >> "
            f"/Contents {content_id} 0 R >>".encode()
        )
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{index} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_csv(r: dict) -> bytes:
    rows = ["PRO Number,Amount Paid"] + [f"{l['pro']},{l['amount']:.2f}" for l in r["lines"]]
    return ("\n".join(rows) + "\n").encode()


def make_xlsx(r: dict) -> bytes | None:
    try:
        from openpyxl import Workbook
    except ImportError:
        return None
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Remittance"
    sheet.append(["PRO Number", "Amount Paid"])
    for line in r["lines"]:
        sheet.append([line["pro"], line["amount"]])
        sheet.cell(row=sheet.max_row, column=1).number_format = "@"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- message builder -----------------------------------------------------------------------------


def build_message(kind: str, rng: random.Random) -> tuple[str, bytes, dict]:
    r = make_remittance(rng)
    domain = r["payer"].split()[0].lower().replace("'", "") + ".example.com"
    message = EmailMessage()
    message["From"] = f"{r['payer']} Accounts Payable <ap@{domain}>"
    message["To"] = "cashapplication@foundryfreight.example.com"
    message["Date"] = format_datetime(datetime.now(timezone.utc))
    message["Message-ID"] = make_msgid(domain=domain)
    meta = {"kind": kind, "payer": r["payer"], "lines": len(r["lines"]), "total": r["total"]}
    footer = f"\n\nThank you,\n{r['payer']} Accounts Payable\nref {r['nonce']}\n"

    if kind == "exception":
        variant = rng.choice(["short_pay", "portal_link", "conflict"])
        meta["variant"] = variant
        if variant == "short_pay":
            stated = round(r["total"] - rng.choice([25.0, 42.35, 118.0, 310.5]), 2)
            message["Subject"] = f"ACH Remittance - {r['payer']} - {r['reference']}"
            message.set_content(
                "Please see the attached remittance. Freight claim deduction taken on one PRO.\n\n"
                + remittance_text(r, total=stated) + footer
            )
            message.add_attachment(make_pdf(remittance_text(r, total=stated).splitlines()), maintype="application",
                                   subtype="pdf", filename=f"Remittance_{r['reference']}.pdf")
        elif variant == "portal_link":
            message["Subject"] = f"Payment notification {r['reference']}"
            message.set_content(
                f"An ACH payment of USD {r['total']:,.2f} was sent on {r['date']}.\n"
                f"View remittance detail: https://payables.{domain}/remit/{r['nonce']}\n" + footer
            )
        else:
            message["Subject"] = f"ACH Remittance {r['reference']}"
            body_lines = dict(r)
            body_lines["lines"] = [dict(l) for l in r["lines"]]
            body_lines["lines"][0]["amount"] = round(body_lines["lines"][0]["amount"] + 50, 2)
            message.set_content(remittance_text(body_lines) + footer)
            message.add_attachment(make_csv(r), maintype="text", subtype="csv", filename=f"remit_{r['reference']}.csv")
    elif kind == "body":
        message["Subject"] = f"ACH Payment Remittance {r['reference']}"
        message.set_content("Remittance detail for today's ACH payment:\n\n" + remittance_text(r) + footer)
    elif kind == "pdf":
        message["Subject"] = f"Remittance Advice - {r['payer']}"
        message.set_content(f"Remittance advice for ACH {r['reference']} attached." + footer)
        message.add_attachment(make_pdf(remittance_text(r).splitlines()), maintype="application", subtype="pdf",
                               filename=f"Remittance_Advice_{r['reference']}.pdf")
    elif kind == "csv":
        message["Subject"] = f"ACH remittance {r['reference']}"
        message.set_content(
            f"Payer: {r['payer']}\nPayment date: {r['date']}\nPayment total: USD {r['total']:,.2f}\n"
            "Line detail attached." + footer
        )
        message.add_attachment(make_csv(r), maintype="text", subtype="csv", filename=f"remittance_{r['reference']}.csv")
    else:  # xlsx
        workbook = make_xlsx(r)
        message["Subject"] = f"ACH remittance detail {r['reference']}"
        message.set_content(
            f"Payer: {r['payer']}\nPayment date: {r['date']}\nPayment total: USD {r['total']:,.2f}\n"
            "Remittance workbook attached." + footer
        )
        if workbook:
            message.add_attachment(workbook, maintype="application",
                                   subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                   filename=f"remittance_{r['reference']}.xlsx")
        else:
            message.add_attachment(make_csv(r), maintype="text", subtype="csv", filename=f"remittance_{r['reference']}.csv")
            meta["kind"] = "csv"

    name = f"{kind}-{r['reference']}-{r['nonce']}.eml"
    return name, message.as_bytes(), meta


# --- submission ----------------------------------------------------------------------------------


def submit(api_base: str, name: str, data: bytes, timeout: float = 60) -> dict:
    boundary = uuid.uuid4().hex
    body = (
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{name}\"\r\n"
        "Content-Type: message/rfc822\r\n\r\n"
    ).encode() + data + f"\r\n--{boundary}--\r\n".encode()
    request = urllib.request.Request(
        f"{api_base.rstrip('/')}/api/remittances?stream=loadtest",
        data=body,
        method="POST",
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        return json.loads(response.read())


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--api", help="Base URL of the demo UI (its nginx proxies /api) or the API.")
    parser.add_argument("--count", type=int, default=25)
    parser.add_argument("--duration-minutes", type=float, default=0, help="Spread submissions evenly over this window (0 = as fast as possible).")
    parser.add_argument("--concurrency", type=int, default=4)
    parser.add_argument("--mix", default=",".join(f"{k}={v}" for k, v in MIX.items()), help="Weights, e.g. body=30,pdf=35,csv=15,xlsx=10,exception=10")
    parser.add_argument("--duplicate-rate", type=float, default=0.0, help="Fraction of messages re-sent byte-identical (exercises dedupe).")
    parser.add_argument("--seed", type=int, default=None)
    parser.add_argument("--out-dir", help="Also write each generated .eml here.")
    parser.add_argument("--dry-run", action="store_true", help="Generate only; do not submit.")
    parser.add_argument("--yes", action="store_true", help="Confirm billable runs above 50 messages.")
    args = parser.parse_args()

    if not args.dry_run and not args.api:
        parser.error("--api is required unless --dry-run")
    if not args.dry_run and args.count > 50 and not args.yes:
        parser.error("runs above 50 messages call billable models; re-run with --yes")

    weights = {}
    for part in args.mix.split(","):
        key, _, value = part.partition("=")
        if key.strip() in MIX:
            weights[key.strip()] = float(value)
    rng = random.Random(args.seed)
    kinds = rng.choices(list(weights), weights=list(weights.values()), k=args.count)
    out_dir = Path(args.out_dir) if args.out_dir else None
    if out_dir:
        out_dir.mkdir(parents=True, exist_ok=True)

    interval = (args.duration_minutes * 60 / args.count) if args.duration_minutes and args.count else 0
    results, lock = [], threading.Lock()
    started = time.perf_counter()
    previous = None

    def task(index: int, kind: str, payload):
        name, data, meta = payload
        if interval:
            delay = started + index * interval - time.perf_counter()
            if delay > 0:
                time.sleep(delay)
        if args.dry_run:
            return {"name": name, **meta, "status": "generated"}
        t0 = time.perf_counter()
        try:
            response = submit(args.api, name, data)
            outcome = {"name": name, **meta, "status": "queued", "processingId": response.get("processingId")}
        except urllib.error.HTTPError as exc:
            outcome = {"name": name, **meta, "status": f"http_{exc.code}", "detail": exc.read()[:200].decode(errors="replace")}
        except Exception as exc:
            outcome = {"name": name, **meta, "status": "error", "detail": type(exc).__name__}
        outcome["submitMs"] = int((time.perf_counter() - t0) * 1000)
        return outcome

    payloads = []
    for kind in kinds:
        if previous and args.duplicate_rate and rng.random() < args.duplicate_rate:
            payloads.append((kind, (f"dup-{previous[0]}", previous[1], {**previous[2], "kind": "duplicate"})))
            continue
        built = build_message(kind, rng)
        previous = built
        payloads.append((kind, built))
        if out_dir:
            (out_dir / built[0]).write_bytes(built[1])

    print(f"[load] {args.count} message(s) | mix {weights} | "
          f"{'pacing ' + format(interval, '.2f') + 's apart' if interval else 'no pacing'} | concurrency {args.concurrency}")
    with ThreadPoolExecutor(max_workers=max(1, args.concurrency)) as pool:
        futures = [pool.submit(task, i, kind, payload) for i, (kind, payload) in enumerate(payloads)]
        for done_count, future in enumerate(as_completed(futures), start=1):
            outcome = future.result()
            with lock:
                results.append(outcome)
            if done_count % 25 == 0 or done_count == args.count:
                elapsed = time.perf_counter() - started
                print(f"[load] {done_count}/{args.count} submitted in {elapsed:.0f}s")

    by_status, by_kind = {}, {}
    for r in results:
        by_status[r["status"]] = by_status.get(r["status"], 0) + 1
        by_kind[r["kind"]] = by_kind.get(r["kind"], 0) + 1
    submit_ms = sorted(r["submitMs"] for r in results if "submitMs" in r)
    print(json.dumps({
        "submitted": len(results),
        "byStatus": by_status,
        "byKind": by_kind,
        "submitP50Ms": submit_ms[len(submit_ms) // 2] if submit_ms else None,
        "elapsedSeconds": round(time.perf_counter() - started, 1),
    }, indent=2))
    if out_dir:
        (out_dir / "load-results.json").write_text(json.dumps(results, indent=2))
    return 0 if all(r["status"] in ("queued", "generated") for r in results) else 1


if __name__ == "__main__":
    sys.exit(main())
